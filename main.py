"""
HUFEM - Hava ve Uzay Fizyolojisi Eğitim Merkezi
QML-based Modern UI Application
"""

import sys
import os
from pathlib import Path
from datetime import datetime
import io

from PySide6.QtCore import QObject, Slot, Signal, Property, QUrl
from PySide6.QtGui import QGuiApplication
from PySide6.QtQml import QQmlApplicationEngine, qmlRegisterSingletonType

# PDF and Excel libraries
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib.enums import TA_CENTER, TA_LEFT
import pandas as pd

# ==================== DATABASE ====================

import sqlite3
import json
from datetime import datetime


class Database:
    """SQLite Database Manager"""
    
    def __init__(self, db_path: str = "hufem.db"):
        self.db_path = db_path
        self.conn = None
        self.connect()
        self.create_tables()
        self.insert_defaults()
    
    def connect(self):
        self.conn = sqlite3.connect(self.db_path)
        self.conn.row_factory = sqlite3.Row
    
    def close(self):
        if self.conn:
            self.conn.close()
    
    def create_tables(self):
        cursor = self.conn.cursor()
        
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS devices (
                id INTEGER PRIMARY KEY,
                name TEXT NOT NULL,
                code TEXT UNIQUE NOT NULL,
                color TEXT
            )
        """)
        
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS participants (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                tc_no TEXT,
                first_name TEXT NOT NULL,
                last_name TEXT NOT NULL,
                title TEXT,
                branch TEXT,
                institution TEXT,
                email TEXT,
                phone TEXT,
                notes TEXT,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS training_groups (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                description TEXT,
                start_date TEXT,
                end_date TEXT,
                status TEXT DEFAULT 'active',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS sessions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                group_id INTEGER,
                device_id INTEGER NOT NULL,
                session_date TEXT NOT NULL,
                supervisor TEXT,
                notes TEXT,
                status TEXT DEFAULT 'completed',
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (group_id) REFERENCES training_groups(id),
                FOREIGN KEY (device_id) REFERENCES devices(id)
            )
        """)
        
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS surveys (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL,
                device_id INTEGER,
                description TEXT,
                questions_json TEXT,
                is_active INTEGER DEFAULT 1,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (device_id) REFERENCES devices(id)
            )
        """)
        
        cursor.execute("""
            CREATE TABLE IF NOT EXISTS survey_responses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                session_id INTEGER,
                survey_id INTEGER,
                participant_id INTEGER,
                responses_json TEXT,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (session_id) REFERENCES sessions(id),
                FOREIGN KEY (survey_id) REFERENCES surveys(id),
                FOREIGN KEY (participant_id) REFERENCES participants(id)
            )
        """)
        
        self.conn.commit()
    
    def insert_defaults(self):
        cursor = self.conn.cursor()
        
        devices = [
            (1, 'Hipoksi Eğitim Cihazı', 'HYPOXIA', '#F59E0B'),
            (2, 'SD Eğitim Cihazı', 'SD', '#8B5CF6'),
            (3, 'Gece Görüş Laboratuvarı', 'NVG', '#10B981'),
        ]
        
        for d in devices:
            cursor.execute("""
                INSERT OR IGNORE INTO devices (id, name, code, color) VALUES (?, ?, ?, ?)
            """, d)
        
        self.conn.commit()
    
    def get_all_participants(self):
        cursor = self.conn.cursor()
        cursor.execute("SELECT * FROM participants ORDER BY last_name, first_name")
        return [dict(row) for row in cursor.fetchall()]
    
    def add_participant(self, data: dict):
        cursor = self.conn.cursor()
        cursor.execute("""
            INSERT INTO participants (tc_no, first_name, last_name, title, branch, institution, email, phone, notes)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            data.get('tc_no'), data.get('first_name'), data.get('last_name'),
            data.get('title'), data.get('branch'), data.get('institution'),
            data.get('email'), data.get('phone'), data.get('notes')
        ))
        self.conn.commit()
        return cursor.lastrowid
    
    def update_participant(self, id: int, data: dict):
        cursor = self.conn.cursor()
        cursor.execute("""
            UPDATE participants SET 
                tc_no=?, first_name=?, last_name=?, title=?, branch=?,
                institution=?, email=?, phone=?, notes=?
            WHERE id=?
        """, (
            data.get('tc_no'), data.get('first_name'), data.get('last_name'),
            data.get('title'), data.get('branch'), data.get('institution'),
            data.get('email'), data.get('phone'), data.get('notes'), id
        ))
        self.conn.commit()
    
    def delete_participant(self, id: int):
        cursor = self.conn.cursor()
        cursor.execute("DELETE FROM participants WHERE id=?", (id,))
        self.conn.commit()
    
    def get_dashboard_stats(self):
        cursor = self.conn.cursor()
        stats = {}
        
        cursor.execute("SELECT COUNT(*) as c FROM participants")
        stats['total_participants'] = cursor.fetchone()['c']
        
        cursor.execute("SELECT COUNT(*) as c FROM training_groups WHERE status='active'")
        stats['active_groups'] = cursor.fetchone()['c']
        
        cursor.execute("SELECT COUNT(*) as c FROM sessions")
        stats['total_sessions'] = cursor.fetchone()['c']
        
        cursor.execute("""
            SELECT d.code, COUNT(s.id) as count
            FROM devices d
            LEFT JOIN sessions s ON d.id = s.device_id
            GROUP BY d.id
        """)
        stats['by_device'] = {row['code']: row['count'] for row in cursor.fetchall()}
        
        return stats


# ==================== QML BACKEND ====================

class Backend(QObject):
    """Backend for QML communication"""
    
    # Signals
    participantsChanged = Signal()
    statsChanged = Signal()
    
    def __init__(self, parent=None):
        super().__init__(parent)
        self.db = Database()
        self._participants = []
        self._stats = {}
        self.refresh()
    
    def refresh(self):
        self._participants = self.db.get_all_participants()
        self._stats = self.db.get_dashboard_stats()
        self.participantsChanged.emit()
        self.statsChanged.emit()
    
    # Properties
    @Property(list, notify=participantsChanged)
    def participants(self):
        return self._participants
    
    @Property('QVariant', notify=statsChanged)
    def stats(self):
        return self._stats
    
    # Slots
    @Slot(str, str, str, str, str, str, result=bool)
    def addParticipant(self, tc_no, first_name, last_name, title, branch, institution):
        try:
            self.db.add_participant({
                'tc_no': tc_no,
                'first_name': first_name,
                'last_name': last_name,
                'title': title,
                'branch': branch,
                'institution': institution
            })
            self.refresh()
            return True
        except Exception as e:
            print(f"Error adding participant: {e}")
            return False
    
    @Slot(int, result=bool)
    def deleteParticipant(self, id):
        try:
            self.db.delete_participant(id)
            self.refresh()
            return True
        except Exception as e:
            print(f"Error deleting participant: {e}")
            return False
    
    @Slot(result=int)
    def totalParticipants(self):
        return self._stats.get('total_participants', 0)
    
    @Slot(result=int)
    def totalSessions(self):
        return self._stats.get('total_sessions', 0)
    
    @Slot(result=int)
    def activeGroups(self):
        return self._stats.get('active_groups', 0)
    
    @Slot(str, result=int)
    def deviceSessionCount(self, device_code):
        return self._stats.get('by_device', {}).get(device_code, 0)
    
    @Slot(str, result=bool)
    def exportParticipantsToPDF(self, filename):
        """Export participants list to PDF"""
        try:
            if not filename.endswith('.pdf'):
                filename += '.pdf'
            
            # Create PDF document
            doc = SimpleDocTemplate(filename, pagesize=A4)
            story = []
            styles = getSampleStyleSheet()
            
            # Custom styles
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor=colors.HexColor('#6366F1'),
                spaceAfter=30,
                alignment=TA_CENTER
            )
            
            # Title
            story.append(Paragraph("HUFEM - Katılımcı Listesi", title_style))
            story.append(Spacer(1, 0.3*inch))
            
            # Subtitle with date
            subtitle_style = ParagraphStyle(
                'Subtitle',
                parent=styles['Normal'],
                fontSize=10,
                textColor=colors.grey,
                alignment=TA_CENTER
            )
            story.append(Paragraph(f"Rapor Tarihi: {datetime.now().strftime('%d.%m.%Y %H:%M')}", subtitle_style))
            story.append(Spacer(1, 0.5*inch))
            
            # Get participants data
            participants = self._participants
            
            # Create table data
            table_data = [['No', 'TC No', 'Ad Soyad', 'Ünvan', 'Branş', 'Kurum']]
            for i, p in enumerate(participants, 1):
                table_data.append([
                    str(i),
                    p.get('tc_no', ''),
                    f"{p.get('first_name', '')} {p.get('last_name', '')}",
                    p.get('title', ''),
                    p.get('branch', ''),
                    p.get('institution', '')
                ])
            
            # Create table
            table = Table(table_data, colWidths=[0.5*inch, 1.2*inch, 2*inch, 1*inch, 1.5*inch, 1.5*inch])
            table.setStyle(TableStyle([
                # Header styling
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6366F1')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                
                # Body styling
                ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
                ('ALIGN', (0, 1), (0, -1), 'CENTER'),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 9),
                ('TOPPADDING', (0, 1), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
                
                # Grid
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F9FAFB')])
            ]))
            
            story.append(table)
            
            # Footer
            story.append(Spacer(1, 0.5*inch))
            footer_style = ParagraphStyle(
                'Footer',
                parent=styles['Normal'],
                fontSize=8,
                textColor=colors.grey,
                alignment=TA_CENTER
            )
            story.append(Paragraph(f"Toplam {len(participants)} katılımcı", footer_style))
            story.append(Paragraph("Sağlık Bilimleri Üniversitesi - HUFEM", footer_style))
            
            # Build PDF
            doc.build(story)
            return True
        except Exception as e:
            print(f"Error exporting to PDF: {e}")
            return False
    
    @Slot(str, result=bool)
    def exportParticipantsToExcel(self, filename):
        """Export participants list to Excel"""
        try:
            if not filename.endswith('.xlsx'):
                filename += '.xlsx'
            
            # Get participants data
            participants = self._participants
            
            # Convert to DataFrame
            df = pd.DataFrame(participants)
            
            # Select and rename columns
            columns_map = {
                'id': 'ID',
                'tc_no': 'TC Kimlik No',
                'first_name': 'Ad',
                'last_name': 'Soyad',
                'title': 'Ünvan',
                'branch': 'Branş',
                'institution': 'Kurum',
                'email': 'E-posta',
                'phone': 'Telefon',
                'created_at': 'Kayıt Tarihi'
            }
            
            df = df[[col for col in columns_map.keys() if col in df.columns]]
            df.rename(columns=columns_map, inplace=True)
            
            # Create Excel writer with formatting
            with pd.ExcelWriter(filename, engine='openpyxl') as writer:
                df.to_excel(writer, sheet_name='Katılımcılar', index=False)
                
                # Get workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Katılımcılar']
                
                # Style the header
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                
                header_fill = PatternFill(start_color='6366F1', end_color='6366F1', fill_type='solid')
                header_font = Font(color='FFFFFF', bold=True, size=11)
                header_alignment = Alignment(horizontal='center', vertical='center')
                
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = header_alignment
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(cell.value)
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Add borders to all cells
                thin_border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
                
                for row in worksheet.iter_rows(min_row=1, max_row=len(df)+1):
                    for cell in row:
                        cell.border = thin_border
                        if cell.row > 1:  # Not header
                            cell.alignment = Alignment(horizontal='left', vertical='center')
            
            return True
        except Exception as e:
            print(f"Error exporting to Excel: {e}")
            return False
    
    @Slot(int, str, result=bool)
    def exportSurveyResultsToPDF(self, survey_id, filename):
        """Export survey results to PDF"""
        try:
            if not filename.endswith('.pdf'):
                filename += '.pdf'
            
            # Create PDF document
            doc = SimpleDocTemplate(filename, pagesize=A4)
            story = []
            styles = getSampleStyleSheet()
            
            # Custom styles
            title_style = ParagraphStyle(
                'CustomTitle',
                parent=styles['Heading1'],
                fontSize=24,
                textColor=colors.HexColor('#6366F1'),
                spaceAfter=30,
                alignment=TA_CENTER
            )
            
            # Title
            story.append(Paragraph("HUFEM - Anket Sonuçları", title_style))
            story.append(Spacer(1, 0.3*inch))
            
            # Subtitle
            subtitle_style = ParagraphStyle(
                'Subtitle',
                parent=styles['Normal'],
                fontSize=10,
                textColor=colors.grey,
                alignment=TA_CENTER
            )
            story.append(Paragraph(f"Rapor Tarihi: {datetime.now().strftime('%d.%m.%Y %H:%M')}", subtitle_style))
            story.append(Spacer(1, 0.5*inch))
            
            # Add survey info and results here
            # This is a placeholder - you would query actual survey data
            info_text = f"Anket ID: {survey_id}"
            story.append(Paragraph(info_text, styles['Normal']))
            story.append(Spacer(1, 0.3*inch))
            
            # Build PDF
            doc.build(story)
            return True
        except Exception as e:
            print(f"Error exporting survey to PDF: {e}")
            return False


# ==================== MAIN ====================

def main():
    # High DPI support
    os.environ["QT_AUTO_SCREEN_SCALE_FACTOR"] = "1"
    
    app = QGuiApplication(sys.argv)
    app.setOrganizationName("Sağlık Bilimleri Üniversitesi")
    app.setOrganizationDomain("sbu.edu.tr")
    app.setApplicationName("HUFEM")
    
    # Create backend
    backend = Backend()
    
    # Create QML engine
    engine = QQmlApplicationEngine()
    
    # Add import paths
    qml_dir = Path(__file__).parent / "qml"
    engine.addImportPath(str(qml_dir))
    
    # Expose backend to QML
    engine.rootContext().setContextProperty("backend", backend)
    
    # Load main QML
    qml_file = qml_dir / "main.qml"
    engine.load(QUrl.fromLocalFile(str(qml_file)))
    
    if not engine.rootObjects():
        print("Error: Could not load QML file")
        sys.exit(-1)
    
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
